"""Shared plumbing for CSV/Excel imports (docs/0084, docs/0085).

Dry run by default: an import runs exactly as it would for real inside the transaction, then
rolls back, so the preview reports everything a real run would (including database errors).
With apply=True it commits. All or nothing: any row error means nothing is written.
"""

import csv
import io
from dataclasses import dataclass, field
from typing import Any

from openpyxl import load_workbook
from pydantic import BaseModel
from sqlalchemy.ext.asyncio import AsyncSession

from hsp import models
from hsp.audit import audit_event
from hsp.auth.principal import Principal
from hsp.problems import ProblemException

MAX_ROWS = 50_000
COUNTS = {"rows", "created", "updated", "unchanged", "reactivated", "positions_created"}


class RowError(BaseModel):
    row: int  # 1-based spreadsheet row; the header is row 1
    column: str | None = None
    code: str
    message: str


class ImportReport(BaseModel):
    kind: str
    applied: bool
    rows: int
    created: int = 0
    updated: int = 0
    unchanged: int = 0
    reactivated: int = 0
    positions_created: int = 0
    errors: list[RowError] = []
    warnings: list[str] = []


def read_table(filename: str, data: bytes) -> list[dict[str, str | None]]:
    """Rows as dicts keyed by lowercase header. CSV (UTF-8, BOM ok) or XLSX (first sheet)."""
    name = filename.lower()
    if name.endswith(".xlsx"):
        sheet = load_workbook(io.BytesIO(data), read_only=True, data_only=True).worksheets[0]
        raw = [
            [("" if c is None else str(c)) for c in r] for r in sheet.iter_rows(values_only=True)
        ]
    elif name.endswith(".csv"):
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError:
            raise ProblemException(
                422, "import-unreadable", "Unreadable file", "CSV files must be UTF-8."
            ) from None
        raw = list(csv.reader(io.StringIO(text)))
    else:
        raise ProblemException(
            415, "import-format", "Unsupported file", "Upload a .csv or .xlsx file."
        )
    if not raw:
        return []
    header = [h.strip().lower() for h in raw[0]]
    rows = []
    for values in raw[1:]:
        if not any(v.strip() for v in values):
            continue  # skip blank lines
        rows.append({h: (v.strip() or None) for h, v in zip(header, values, strict=False) if h})
    if len(rows) > MAX_ROWS:
        raise ProblemException(
            413, "import-too-large", "Too many rows", f"At most {MAX_ROWS} rows per import."
        )
    return rows


@dataclass
class ImportRun:
    report: ImportReport
    row_of: dict[str, int] = field(default_factory=dict)

    def error(self, row: int, code: str, message: str, column: str | None = None) -> None:
        self.report.errors.append(RowError(row=row, column=column, code=code, message=message))


def check_columns(
    rows: list[dict[str, Any]], known: set[str], required: set[str], run: ImportRun
) -> bool:
    present = set(rows[0]) if rows else set()
    missing = required - present
    if rows and missing:
        run.error(1, "missing_column", f"Missing required column(s): {', '.join(sorted(missing))}.")
        return False
    unknown = present - known
    if unknown:
        run.report.warnings.append(f"Ignored unknown column(s): {', '.join(sorted(unknown))}.")
    return True


async def finish_import(
    db: AsyncSession,
    principal: Principal,
    request_id: str | None,
    run: ImportRun,
    apply: bool,
) -> ImportReport:
    """Roll back a dry run or a run with errors; otherwise audit the counts and commit."""
    if run.report.errors or not apply:
        await db.rollback()
        run.report.applied = False
        return run.report
    db.add(
        audit_event(
            principal,
            action=f"import.{run.report.kind}",
            entity_type="import",
            entity_id=models.new_id(),
            after=run.report.model_dump(include=COUNTS),
            request_id=request_id,
        )
    )
    await db.commit()
    run.report.applied = True
    return run.report
